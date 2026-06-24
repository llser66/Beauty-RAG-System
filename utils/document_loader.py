"""
多源数据资产引擎 — 多格式文档解析器
支持 .txt .pdf .docx .xlsx 自动扫描与智能切片
"""

import os
import glob
from typing import List, Dict, Any
from langchain_core.documents import Document


def load_single_txt(filepath: str, encoding: str = "utf-8") -> List[Document]:
    """加载单个 .txt 文件，按段落返回"""
    with open(filepath, "r", encoding=encoding) as f:
        text = f.read()
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    docs = []
    for para in paragraphs:
        docs.append(Document(
            page_content=para,
            metadata={"source": filepath, "type": "txt"}
        ))
    return docs


def load_single_pdf(filepath: str) -> List[Document]:
    """加载单个 .pdf 文件，按页面提取"""
    import fitz  # PyMuPDF
    docs = []
    with fitz.open(filepath) as pdf:
        for page_num, page in enumerate(pdf):
            text = page.get_text().strip()
            if text:
                docs.append(Document(
                    page_content=text,
                    metadata={"source": filepath, "page": page_num + 1, "type": "pdf"}
                ))
    return docs


def load_single_docx(filepath: str) -> List[Document]:
    """加载单个 .docx 文件，按段落提取"""
    from docx import Document as DocxDocument
    docs = []
    doc = DocxDocument(filepath)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for i, para in enumerate(paragraphs):
        docs.append(Document(
            page_content=para,
            metadata={"source": filepath, "para": i, "type": "docx"}
        ))
    return docs


def load_single_xlsx(filepath: str) -> List[Document]:
    """加载单个 .xlsx 文件，每行作为一个文档，保留表头"""
    import openpyxl
    docs = []
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h) if h else f"col_{i}" for i, h in enumerate(rows[0])]
        for row_idx, row in enumerate(rows[1:], start=2):
            content = " | ".join(
                f"{headers[i]}: {str(cell)}" for i, cell in enumerate(row) if cell is not None
            )
            if content.strip():
                docs.append(Document(
                    page_content=content,
                    metadata={
                        "source": filepath, "sheet": sheet_name,
                        "row": row_idx, "type": "xlsx"
                    }
                ))
    wb.close()
    return docs


def load_single_file(filepath: str) -> List[Document]:
    """
    根据文件扩展名自动选择解析器，加载单个文件
    返回 Document 列表（未经切片）
    """
    ext = os.path.splitext(filepath)[1].lower()
    loader = LOADER_MAP.get(ext)
    if loader is None:
        raise ValueError(f"不支持的文件格式: {ext}")
    print(f"  📄 加载: {filepath}")
    docs = loader(filepath)
    print(f"     → 提取 {len(docs)} 个片段")
    return docs


LOADER_MAP = {
    ".txt": load_single_txt,
    ".pdf": load_single_pdf,
    ".docx": load_single_docx,
    ".xlsx": load_single_xlsx,
}


def scan_knowledge_base(directory: str) -> List[Document]:
    """
    扫描 knowledge_base/ 目录下所有支持格式的文档，返回 Document 列表
    """
    all_docs = []
    supported = "|".join(LOADER_MAP.keys())
    pattern = os.path.join(directory, f"*{{{supported}}}") if "," in supported else os.path.join(directory, f"*{supported}")

    for ext, loader in LOADER_MAP.items():
        for filepath in glob.glob(os.path.join(directory, f"*{ext}")):
            print(f"  📄 加载: {filepath}")
            try:
                docs = loader(filepath)
                all_docs.extend(docs)
                print(f"     → 提取 {len(docs)} 个切片")
            except Exception as e:
                print(f"     ❌ 解析失败: {e}")
    return all_docs


def scan_knowledge_base_with_txt(directory: str, extra_txt_path: str = None) -> List[Document]:
    """
    扫描目录 + 可选附加单个 .txt 文件
    """
    docs = scan_knowledge_base(directory)
    if extra_txt_path and os.path.exists(extra_txt_path):
        try:
            extra = load_single_txt(extra_txt_path)
            docs.extend(extra)
            print(f"  📄 加载附加文件: {extra_txt_path} → {len(extra)} 个切片")
        except Exception as e:
            print(f"     ❌ 附加文件解析失败: {e}")
    return docs
